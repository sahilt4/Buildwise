from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'documents' / 'Buildwise_PBL_Project_Report.docx'
ARCH = ROOT / 'output' / 'images' / 'buildwise_architecture_supplied.jpg'
OUT.parent.mkdir(parents=True, exist_ok=True)
NAVY = '18355F'
LIGHT = 'F4F7FB'

def set_cell(cell, text, header=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.font.name = 'Aptos'
    r.font.size = Pt(8.2)
    r.bold = header
    r.font.color.rgb = RGBColor(255, 255, 255) if header else RGBColor(24, 34, 48)

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), NAVY)
        tcPr.append(shd)
        set_cell(c, h, True)
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            if index % 2:
                tcPr = cells[i]._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), LIGHT)
                tcPr.append(shd)
            set_cell(cells[i], value)
    doc.add_paragraph()

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0, 0, 0)

def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    r.font.name = 'Aptos'
    r.font.size = Pt(10.5)

def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)

def placeholder(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = RGBColor(102, 112, 133)

def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Buildwise PBL Project Report | Department of Computer Engineering | Page ')
    r.font.size = Pt(8)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.7)
section.right_margin = Inches(.7)
add_footer(section)
for name, size in [('Normal', 10.5), ('Title', 22), ('Heading 1', 15), ('Heading 2', 12), ('Heading 3', 11)]:
    style = doc.styles[name]
    style.font.name = 'Aptos'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)

# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('K.K Wagh Institute of Engineering Education and Research, Nashik').bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Department of Computer Engineering').bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(42)
r = p.add_run('PBL PROJECT REPORT')
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = RGBColor(24, 53, 95)
for line in ['Project Based Learning', 'Course Code: 2301310', 'Academic Year: 2026-27', 'Semester: V', 'Class: T.Y. B.Tech. Computer Engineering', 'Division: A', 'PID: ____________________']:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(line)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(20)
r = p.add_run('BUILDWISE')
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = RGBColor(24, 53, 95)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Construction Management and Leftover Material Marketplace').bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
p.add_run('Submitted by').bold = True
add_table(doc, ['Roll No.', 'Student Name', 'CRN', 'Signature'], [['To be filled', 'To be filled', 'To be filled', ''], ['To be filled', 'To be filled', 'To be filled', ''], ['To be filled', 'To be filled', 'To be filled', '']])
placeholder(doc, 'Faculty/Mentor Name: ______________________________')
doc.add_page_break()

heading(doc, 'CERTIFICATE')
body(doc, 'This is to certify that the Project Based Learning project titled “Buildwise - Construction Management and Leftover Material Marketplace” has been completed as a submitted prototype for Course Code 2301310 during Academic Year 2026-27. The project demonstrates problem identification, system design, implementation, testing, documentation and presentation through a React-based client-side application.')
placeholder(doc, 'Faculty/Mentor Signature: ____________________     Date: ____________________')
heading(doc, 'STUDENT DECLARATION')
body(doc, 'We declare that this report describes the Buildwise prototype developed by the student team as part of Project Based Learning. The source, design decisions, prototype evidence and references are represented honestly.')
placeholder(doc, 'Student signatures: ______________________________')
heading(doc, 'ACKNOWLEDGEMENT')
body(doc, 'The team acknowledges the faculty mentor, the Department of Computer Engineering and K.K Wagh Institute of Engineering Education and Research. The React, Vite, jsPDF, Lucide and open-source communities are acknowledged for documentation and tooling that supported the prototype implementation.')
heading(doc, 'ABSTRACT')
body(doc, 'Buildwise is a client-side construction operations prototype that combines multi-site oversight, material inventory, workforce attendance, task coordination and a leftover-material marketplace. Construction teams often maintain material balances, attendance and task progress across separate registers or spreadsheets, making stock visibility and surplus recovery difficult. The proposed solution provides a responsive React interface for builders, engineers and workers. It tracks sites, materials, leftovers, marketplace listings, workers, tasks, activities and attendance in React Context state, with browser localStorage used to retain prototype data after refresh. Material usage updates remaining quantity, identifies low stock using thresholds and can detect recoverable surplus. Leftover records can be listed in a marketplace, while workers can use one-click attendance or a simulated QR verification flow. The application also produces client-side PDF reports and CSV exports. The current implementation is intentionally client-only: authentication, real camera scanning, backend synchronization, payment workflows and cloud databases are not implemented.')
body(doc, 'Keywords: Construction management, material inventory, workforce attendance, React, circular economy, marketplace')
doc.add_page_break()

heading(doc, 'TABLE OF CONTENTS')
body(doc, '1. Introduction\n2. Problem Identification and Background Study\n3. Problem Definition and Requirements\n4. Project Planning and Gantt Chart\n5. System Design and Architecture\n6. Module Identification and Technology Selection\n7. Implementation\n8. Testing and Validation\n9. Results and Discussion\n10. Project Demonstration\n11. Ethical, Social and Sustainability Considerations\n12. Teamwork and Individual Contributions\n13. Challenges Faced and Solutions Adopted\n14. Learning Outcomes and Reflection\n15. Conclusion and Future Scope\n16. References\n17. Appendices')
doc.add_page_break()

heading(doc, '1. INTRODUCTION')
heading(doc, '1.1 Background', 2)
body(doc, 'Construction projects require coordinated visibility of project sites, material stock, workforce attendance and task progress. Buildwise is a browser-based prototype for these operational needs, with an additional workflow that turns eligible leftover material into a marketplace listing.')
heading(doc, '1.2 Motivation', 2)
body(doc, 'The project addresses fragmented operational records and the loss of potentially reusable construction materials. One prototype interface shows the relationship between material consumption, surplus identification and resale recovery.')
heading(doc, '1.3 Need for the Project', 2)
body(doc, 'The implemented system supports users who need a simple shared view of site status, inventory, workers, attendance and tasks. Its material workflow exposes low-stock risk and surplus candidates before they are treated as waste.')
heading(doc, '1.4 Scope of the Project', 2)
body(doc, 'Included: responsive views for sites, materials, marketplace, workers, attendance, tasks, reports, notifications and settings; builder, engineer and worker perspectives; localStorage persistence; simulated QR scanning; PDF and CSV exports. Excluded: backend APIs, cloud synchronization, real authentication, real camera QR decoding, online payments and verified logistics.')
heading(doc, '1.5 Project Objectives', 2)
bullets(doc, ['Provide a multi-site dashboard for operational visibility.', 'Track material purchase, usage, balance and low-stock status.', 'Convert eligible leftover material into marketplace listings.', 'Record attendance through one-click punch and simulated QR verification.', 'Coordinate tasks and export selected operations data as PDF or CSV.'])

heading(doc, '2. PROBLEM IDENTIFICATION AND BACKGROUND STUDY')
heading(doc, '2.1 Problem Identification', 2)
body(doc, 'Construction teams can lose visibility when site progress, materials, workforce attendance and tasks are tracked independently. This makes it harder to recognize low stock, assign work and recover value from usable leftovers.')
heading(doc, '2.2 Existing System / Existing Solutions', 2)
body(doc, 'The baseline approach is manual registers and separate spreadsheets for stock, attendance and task tracking. Generic project-management tools may organize tasks but do not necessarily connect construction inventory to leftover-material recovery.')
heading(doc, '2.3 Limitations of Existing Solutions', 2)
body(doc, 'Separate records duplicate effort, delay stock decisions and make surplus difficult to identify. Manual hand-offs can result in usable stock being ignored rather than made available for reuse or resale.')
heading(doc, '2.4 Background / Literature Study', 2)
body(doc, 'The implementation uses React and Vite documentation for the interface, jsPDF and AutoTable documentation for browser-side reporting, and UN SDG descriptions for SDG 9 and SDG 12. No external dataset is used; the prototype uses seed records in src/data/mockData.js.')
heading(doc, '2.5 Problem Analysis', 2)
body(doc, 'The project models sites, materials, leftovers, marketplace listings, workers, tasks, activities and attendance logs as connected state entities. AppContext actions update these entities and persist them in localStorage.')
heading(doc, '2.6 SDG / Social Relevance', 2)
add_table(doc, ['Relevant SDG', 'Justification'], [['SDG 9 - Industry, Innovation and Infrastructure', 'Buildwise digitizes connected construction workflows for sites, materials, workforce, tasks and reporting.'], ['SDG 12 - Responsible Consumption and Production', 'Material tracking, surplus identification and marketplace listing support reuse and reduce the likelihood of reusable construction material being discarded.']])
doc.add_page_break()

heading(doc, '3. PROBLEM DEFINITION AND REQUIREMENTS')
heading(doc, '3.1 Problem Statement', 2)
body(doc, 'Develop a responsive, client-side construction management prototype that helps builders, engineers and workers monitor sites, materials, attendance and tasks, while allowing eligible leftover construction materials to be listed in a circular marketplace.')
heading(doc, '3.2 Functional Requirements', 2)
add_table(doc, ['Sr. No.', 'Functional Requirement', 'Description'], [('FR-01', 'Role switching', 'Builder, engineer and worker views.'), ('FR-02', 'Dashboard KPIs', 'Active sites, workers, low stock and task indicators.'), ('FR-03', 'Site management', 'Site progress, budget, health and supervisor records.'), ('FR-04', 'Material inventory', 'Add, filter, view and record material usage.'), ('FR-05', 'Stock logic', 'Calculate remaining quantity and low-stock status.'), ('FR-06', 'Surplus recovery', 'Identify eligible surplus and create leftovers.'), ('FR-07', 'Marketplace', 'Post, filter and request surplus listings.'), ('FR-08', 'Worker management', 'Add and browse workers by site/trade.'), ('FR-09', 'Attendance', 'One-click punch, simulated QR scan and muster log.'), ('FR-10', 'Task management', 'Create, filter and update task progress with prerequisites.'), ('FR-11', 'Reports', 'Generate client-side PDF reports and CSV exports.'), ('FR-12', 'Notifications', 'Show activity and low-stock information.')])
heading(doc, '3.3 Non-Functional Requirements', 2)
add_table(doc, ['Requirement', 'Description'], [('Usability', 'Responsive navigation and role-focused views use cards, badges, modals and feedback.'), ('Reliability', 'State collections are saved to browser localStorage after changes.'), ('Maintainability', 'Source is organized into React components, AppContext, mock data, utilities and CSS.'), ('Compatibility', 'Prototype is intended for current evergreen desktop and mobile browsers.'), ('Security boundary', 'React renders user text through JSX; authentication and server authorization are not implemented.'), ('Performance boundary', 'Client-side data operations do not require a network API; the production bundle is generated by Vite.')])
heading(doc, '3.4 Users / Stakeholders', 2)
add_table(doc, ['Stakeholder', 'Requirement / Role'], [('Builder', 'Portfolio overview, sites, materials, marketplace, reports and operational visibility.'), ('Site Engineer', 'Material and task-oriented site management perspective.'), ('Worker', 'Simplified worker portal, attendance punch and assigned-task review.'), ('Academic Evaluator', 'Evidence of requirements, design, implementation and validation.')])

heading(doc, '4. PROJECT PLANNING AND GANTT CHART')
heading(doc, '4.1 Project Team Formation', 2)
add_table(doc, ['Roll No.', 'Student Name', 'Assigned Role'], [('To be filled', 'To be filled', 'To be filled by team'), ('To be filled', 'To be filled', 'To be filled by team'), ('To be filled', 'To be filled', 'To be filled by team')])
heading(doc, '4.2 Work Breakdown Structure', 2)
add_table(doc, ['Task', 'Evidence in submitted project', 'Status'], [('Problem identification', 'Construction operations and leftover-material recovery problem definition.', 'Completed'), ('Requirement analysis', 'SRS and application feature mapping.', 'Completed'), ('System design', 'React component architecture, AppContext and localStorage design.', 'Completed'), ('Implementation', 'Vite/React source in src/.', 'Completed'), ('Testing and validation', 'Production build, lint output and scenario checks.', 'Completed'), ('Documentation', 'SRS, submission PDFs and this report.', 'Completed')])
heading(doc, '4.3 Gantt Chart', 2)
body(doc, 'A dated team schedule was not included in the repository. The completed workflow followed this sequence: problem identification → requirements and design → interface and state implementation → feature validation → documentation and submission preparation. The student team should enter the actual dates before formal submission.')
doc.add_page_break()

heading(doc, '5. SYSTEM DESIGN AND ARCHITECTURE')
heading(doc, '5.1 Proposed System', 2)
body(doc, 'Buildwise is a React single-page application. The App component selects the active view, while AppContext centralizes operational state and action handlers. Each collection initializes from localStorage or seeded mock data and is written back after state changes.')
heading(doc, '5.2 System Architecture', 2)
if ARCH.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ARCH), width=Inches(6.3))
body(doc, 'Figure 5.1: Buildwise system architecture. The users, React/Vite frontend, AppContext business logic, client-side services and browser localStorage represent the implemented prototype. The cloud backend, database, authentication and AI items shown in the right-hand “Future / Planned” panel are roadmap items only and are not implemented in the submitted project.')
heading(doc, '5.3 Data Flow / Workflow', 2)
body(doc, 'Material workflow: add or select material → record usage → calculate remaining quantity → set low-stock or surplus-candidate status → create leftover record when eligible → publish leftover to marketplace. Attendance workflow: worker punch or simulated QR scan → update worker status → append attendance log → persist and display feedback. Task workflow: create task → validate prerequisite for in-progress status → update task status and progress.')
heading(doc, '5.4 Use Case Diagram', 2)
body(doc, 'Primary use cases are role selection, dashboard review, site and material management, usage logging, surplus listing, marketplace inquiry, worker management, attendance recording, task updates, report generation and CSV export. These use cases are implemented as React views, modals and AppContext actions.')
heading(doc, '5.5 Database / Data Design', 2)
body(doc, 'The prototype does not use a server database. It stores sites, materials, leftovers, marketplace listings, workers, tasks, activity records and attendance logs in browser localStorage with bw_ key names. This is a prototype persistence mechanism, not a multi-user database.')
heading(doc, '5.6 Module Description', 2)
add_table(doc, ['Module', 'Description', 'Responsible Member'], [('Dashboard', 'KPI cards, charts and activity overview.', 'To be assigned'), ('Materials', 'Inventory, usage, low-stock and surplus logic.', 'To be assigned'), ('Marketplace', 'Surplus listing, filtering and inquiry feedback.', 'To be assigned'), ('Workers and Attendance', 'Worker directory, punch card, QR simulation and logs.', 'To be assigned'), ('Tasks and Reports', 'Task progress, prerequisite validation, PDF/CSV output.', 'To be assigned')])

heading(doc, '6. MODULE IDENTIFICATION AND TECHNOLOGY SELECTION')
heading(doc, '6.1 Technology Stack', 2)
add_table(doc, ['Component', 'Technology / Tool', 'Reason for Selection'], [('Front End', 'React 19 and React DOM', 'Component-based views and state-driven rendering.'), ('Build Framework', 'Vite 8', 'Local development and production build workflow.'), ('State / Storage', 'React Context and localStorage', 'Centralized prototype state with refresh persistence.'), ('Reporting', 'jsPDF and jspdf-autotable', 'Client-side PDF tables and reports.'), ('UI assets', 'lucide-react, CSS', 'Icons and reusable visual system.'), ('Feedback', 'canvas-confetti', 'Visual feedback for selected success events.'), ('Testing / quality', 'Vite build and Oxlint', 'Build and static-analysis checks.')])
heading(doc, '6.2 Technology Selection Justification', 2)
body(doc, 'React and Context match the need for connected interfaces without a backend. Vite provides the development and bundle workflow. localStorage makes the academic prototype demonstrable after refresh without server setup. jsPDF and AutoTable implement exports already exposed in the application. No database, Node/Express API or authentication service is presented as implemented.')
heading(doc, '6.3 Prototype', 2)
body(doc, 'The prototype contains dashboards, material inventory, leftover marketplace, workers, attendance, tasks, notifications, settings, reports, a landing page and a worker portal. These are implemented as source components in src/components and are selected by App.jsx according to the active view and user role.')
doc.add_page_break()

heading(doc, '7. IMPLEMENTATION')
heading(doc, '7.1 Development Environment', 2)
add_table(doc, ['Area', 'Implemented Environment'], [('Language', 'JavaScript / JSX and CSS'), ('Runtime', 'Browser-based React single-page application'), ('Build tool', 'Vite 8'), ('Key libraries', 'React, React DOM, jsPDF, jspdf-autotable, lucide-react, canvas-confetti'), ('Persistence', 'Browser localStorage'), ('Source organization', 'src/components, src/context, src/data, src/utils and src/styles')])
heading(doc, '7.2 Module-wise Implementation', 2)
for title, text in [('Module 1 - Materials and Marketplace', 'AppContext.recordUsage updates balance and status. Eligible excess material becomes a leftover record. listLeftoverOnMarketplace transfers a leftover into marketplace state.'), ('Module 2 - Workers and Attendance', 'toggleWorkerAttendance changes check-in state and adds an attendance record. scanQRAttendance simulates verification and logs attendance.'), ('Module 3 - Tasks and Reports', 'updateTaskStatus prevents an in-progress update when a prerequisite is incomplete. pdfGenerator.js creates attendance, material-audit and site-summary PDFs, and exportCSV creates browser downloads.')]:
    heading(doc, title, 3)
    body(doc, text)
heading(doc, '7.3 Important Algorithms / Methods', 2)
bullets(doc, ['Remaining stock = purchased quantity - used quantity, limited to a minimum of zero.', 'Low-stock status applies when remaining quantity is at or below the configured threshold.', 'Surplus requires remaining quantity above required quantity and a non-Scrap condition.', 'Estimated surplus value uses quantity, cost per unit and a condition factor.', 'Task progression checks prerequisite completion before allowing In Progress status.'])
heading(doc, '7.4 Source Code Repository', 2)
placeholder(doc, 'Repository Link: To be added by the student team after publishing the repository.\nBranch/Version: Local submission workspace; branch name to be confirmed by the team.')

heading(doc, '8. TESTING AND VALIDATION')
heading(doc, '8.1 Testing Strategy', 2)
body(doc, 'Validation focused on build verification, static analysis and scenario-based checks of implemented state actions. The Vite production build completed successfully. Oxlint reports warnings in the existing codebase but does not report syntax errors that prevent lint completion. No automated unit-test suite is included, so no unit-test coverage claim is made.')
heading(doc, '8.2 Test Cases', 2)
add_table(doc, ['ID', 'Module', 'Input / Action', 'Expected Output', 'Status'], [('TC01', 'Materials', 'Record valid material usage', 'Used and remaining quantities update; feedback appears.', 'Pass'), ('TC02', 'Materials', 'Record usage above remaining quantity', 'Usage modal validation prevents invalid submission.', 'Pass'), ('TC03', 'Marketplace', 'List a leftover', 'Item moves from leftovers to marketplace.', 'Pass'), ('TC04', 'Attendance', 'Punch or simulate QR scan', 'Worker state and attendance log update.', 'Pass'), ('TC05', 'Tasks', 'Move blocked task to In Progress', 'Prerequisite check prevents update.', 'Pass'), ('TC06', 'Reports', 'Download report or export CSV', 'Browser download is initiated.', 'Pass')])
heading(doc, '8.3 Bugs / Issues Identified', 2)
add_table(doc, ['Issue', 'Cause', 'Resolution / Status'], [('No multi-device synchronization', 'Prototype uses localStorage only.', 'Known limitation; future backend scope.'), ('QR scanner does not use hardware', 'The interface intentionally simulates scanning.', 'Known limitation; future camera integration scope.'), ('Lint warnings remain', 'Existing unused imports/parameters and advisories.', 'Warnings identified; no syntax-blocking error.')])
heading(doc, '8.4 Validation', 2)
body(doc, 'The prototype was validated against implemented requirements through the described user flows, source inspection of AppContext actions and component views, a successful Vite production build and static linting. Evidence includes the interface, source code, architecture, requirements tables and generated report exports.')
doc.add_page_break()

heading(doc, '9. RESULTS AND DISCUSSION')
heading(doc, '9.1 Final System', 2)
body(doc, 'The completed Buildwise prototype provides one browser interface for site overview, inventory control, workforce records, attendance, task coordination, leftover-material listings and client-side reports. Core data remains on the current browser and is seeded from mockData.js on first use.')
heading(doc, '9.2 Results', 2)
add_table(doc, ['Area', 'Implemented Result'], [('Operational overview', 'Dashboard surfaces sites, worker presence, low-stock items, tasks and activity.'), ('Inventory control', 'Material quantities and thresholds drive remaining-balance and stock-status feedback.'), ('Circular recovery', 'Eligible leftovers can be moved into a marketplace listing workflow.'), ('Workforce operations', 'Worker records, attendance logs and task assignments are available in role-aware views.'), ('Reporting', 'The prototype can create PDF reports and CSV exports locally.')])
heading(doc, '9.3 Comparison with Existing System', 2)
add_table(doc, ['Parameter', 'Manual / Fragmented Approach', 'Buildwise Prototype'], [('Inventory status', 'Separate registers or spreadsheets', 'Central materials collection with usage and threshold status.'), ('Attendance', 'Manual roll call or separate record', 'Punch/QR simulation creates attendance log.'), ('Leftovers', 'May remain undocumented', 'Leftovers can be identified and listed in marketplace state.'), ('Reporting', 'Manual consolidation', 'Client-side PDF/CSV export utilities.')])
heading(doc, '9.4 Achievement of Objectives', 2)
add_table(doc, ['Objective', 'Achieved?', 'Evidence'], [('Site and operations overview', 'Yes', 'Dashboard, sites, notifications and role views.'), ('Material monitoring and stock status', 'Yes', 'Materials view and AppContext usage logic.'), ('Leftover-material recovery workflow', 'Yes', 'Leftovers and marketplace actions.'), ('Attendance and task coordination', 'Yes', 'Punch/QR simulation, logs and task status logic.'), ('Cloud multi-user operations', 'No - outside scope', 'No backend, authentication or shared database implemented.')])

heading(doc, '10. PROJECT DEMONSTRATION')
heading(doc, '10.1 Demonstration Flow', 2)
body(doc, '1. Launch Buildwise and choose a builder or worker perspective. 2. Review the dashboard. 3. Open Materials, record usage and observe the updated balance/status. 4. Review leftovers and publish one to Marketplace. 5. Open Attendance and perform a punch or simulated QR scan. 6. Create or update a task, including a prerequisite-dependent task. 7. Open Reports or Attendance to download a PDF/CSV export.')
heading(doc, '10.2 Major Features Demonstrated', 2)
bullets(doc, ['Role-oriented views: Builder, Engineer and Worker.', 'Multi-site and material inventory management.', 'Low-stock and surplus candidate identification.', 'Leftover-to-marketplace conversion and inquiry feedback.', 'Worker management, attendance punch and simulated QR scan.', 'Task filtering, status updates and prerequisite validation.', 'Client-side PDF reports and CSV exports.'])
heading(doc, '10.3 Final Screenshots', 2)
body(doc, 'Representative dashboard, material inventory and marketplace screens are included in the Buildwise Prototype Screenshots submission PDF. They demonstrate the primary builder-facing workflows described in this report.')
heading(doc, '10.4 Presentation / Demonstration Evidence', 2)
placeholder(doc, 'Add presentation slides, final demonstration photographs/screenshots, repository URL and mentor/faculty feedback when available.')

heading(doc, '11. ETHICAL, SOCIAL AND SUSTAINABILITY CONSIDERATIONS')
heading(doc, '11.1 Ethical Considerations', 2)
body(doc, 'The prototype stores data only in the browser and does not implement account authentication or server-side access control. Personal-looking seed data is demo content. A production version would need authenticated roles, secure storage, authorization, audit controls and consent-aware handling of worker data. The application does not implement AI decision-making.')
heading(doc, '11.2 Social Impact', 2)
body(doc, 'Buildwise can make site operations easier to inspect for builders, engineers and workers by keeping attendance, tasks and material status visible in one interface. The worker portal reduces the number of steps needed for attendance and task review.')
heading(doc, '11.3 Sustainability', 2)
body(doc, 'The surplus workflow encourages teams to identify reusable materials and list them for reuse or resale. The prototype does not calculate carbon emissions or verified waste diversion, so it should not be interpreted as a measured environmental-impact system.')
heading(doc, '11.4 SDG Mapping', 2)
add_table(doc, ['SDG', 'Project Contribution'], [('SDG 9 - Industry, Innovation and Infrastructure', 'Digitizes connected construction workflows for sites, materials, workforce, tasks and reporting.'), ('SDG 12 - Responsible Consumption and Production', 'Supports identification, reuse and resale of eligible leftover construction material through the marketplace workflow.')])
doc.add_page_break()

heading(doc, '12. TEAMWORK AND INDIVIDUAL CONTRIBUTIONS')
add_table(doc, ['Student', 'Role', 'Major Contribution', 'Evidence'], [('To be filled', 'To be filled', 'To be filled by team member', 'Commit history / screenshots / review'), ('To be filled', 'To be filled', 'To be filled by team member', 'Commit history / screenshots / review'), ('To be filled', 'To be filled', 'To be filled by team member', 'Commit history / screenshots / review')])
heading(doc, '12.1 Team Collaboration', 2)
body(doc, 'The repository does not identify individual contributors or task allocation. The team should record its actual coordination method, meeting cadence and review process here before submitting the report.')
heading(doc, '12.2 Tools Used for Collaboration', 2)
placeholder(doc, 'Tools Used: To be completed by the team. Do not claim GitHub, Drive, Trello, Jira, Teams or AI collaboration tools unless they were actually used.')
heading(doc, '13. CHALLENGES FACED AND SOLUTIONS ADOPTED')
add_table(doc, ['Challenge', 'Impact', 'Solution Adopted', 'Learning'], [('Connected operational state', 'One action affects related entities.', 'Centralized collections and actions in AppContext.', 'Shared state makes workflow relationships explicit.'), ('Prototype persistence', 'Data should remain after refresh.', 'localStorage synchronization for core collections.', 'Client persistence does not replace multi-user storage.'), ('Task dependencies', 'Tasks should respect prerequisites.', 'Status update checks prerequisite task state.', 'Business rules belong near state transitions.'), ('Reporting', 'Operational data needs readable exports.', 'jsPDF, AutoTable and CSV utility.', 'Reports can be generated client-side for prototype scope.')])
heading(doc, '14. LEARNING OUTCOMES AND REFLECTION')
heading(doc, '14.1 What Did We Learn?', 2)
body(doc, 'The project demonstrates how a domain can be represented as connected front-end state entities and workflows. It also demonstrates the importance of separating implemented functionality from planned future features.')
heading(doc, '14.2 Technical Skills Developed', 2)
bullets(doc, ['React functional components, hooks and Context.', 'Vite development and production bundling.', 'Browser localStorage persistence.', 'Client-side PDF and CSV generation.', 'Responsive component and CSS organization.'])
heading(doc, '14.3 Teamwork / Collaboration Skills', 2)
body(doc, 'Individual collaboration details are not available in the repository and must be completed by the student team using only their actual working process.')
heading(doc, '14.4 Problem-Solving Skills', 2)
body(doc, 'The project required converting construction workflows into client-side state transitions: calculating balances, validating material usage, gating task status with prerequisites and transferring leftovers into marketplace listings.')
heading(doc, '14.5 Individual Reflection', 2)
placeholder(doc, 'Student 1 - Name and reflection: To be written by the student.\nStudent 2 - Name and reflection: To be written by the student.\nStudent 3 - Name and reflection: To be written by the student.\nStudent 4 - Name and reflection: To be written by the student.')

heading(doc, '15. CONCLUSION AND FUTURE SCOPE')
heading(doc, '15.1 Conclusion', 2)
body(doc, 'Buildwise delivers a functioning client-side prototype for construction operations management and leftover-material recovery. It connects site, material, workforce, attendance, task and reporting workflows in a structured React application. The marketplace workflow moves eligible leftovers toward reuse or resale.')
heading(doc, '15.2 Future Scope', 2)
bullets(doc, ['Secure backend API, authentication and role authorization.', 'Cloud database and multi-device synchronization.', 'Real camera-based QR validation rather than simulated scanning.', 'Verified marketplace communication, logistics and payment workflows.', 'Measured waste diversion and sustainability analytics.'])
heading(doc, '16. REFERENCES')
bullets(doc, ['React Documentation, https://react.dev/', 'Vite Documentation, https://vite.dev/', 'jsPDF Documentation, https://github.com/parallax/jsPDF', 'jsPDF-AutoTable Documentation, https://github.com/simonbengtsson/jsPDF-AutoTable', 'United Nations Sustainable Development Goals, https://sdgs.un.org/goals', 'Project source: src/App.jsx, src/context/AppContext.jsx, src/data/mockData.js, src/utils/pdfGenerator.js and src/components/.'])
heading(doc, '17. APPENDICES')
body(doc, 'Appendix A: Problem definition and requirements are documented in Sections 2 and 3.\nAppendix B: Architecture and data design are documented in Section 5.\nAppendix C: Test cases and issue log are documented in Section 8.\nAppendix D: Prototype screenshot evidence is available in the Buildwise Prototype Screenshots submission PDF.\nAppendix E: Team-specific evidence, dates, repository URL, mentor feedback and individual reflections must be added by the student team before submission.')

doc.save(OUT)
print(OUT)
